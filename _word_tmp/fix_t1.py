# -*- coding: utf-8 -*-
"""修复审核表 T1：合并单元格网格定位 + 段落级重建（防 run 拆分问题）"""
import copy
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

T1 = r"C:\Users\郭耀彬\Desktop\gybct\ai-chat-graduation\3.2301132013-XXX-题目审核表（填写版）.docx"
PLACEHOLDER = "×××"

def para_text(p):
    return "".join(r.text or "" for r in p.runs)

def rebuild_para(p, text):
    """清空段落所有 run，复制首个 run 的 rPr 重建单个 run"""
    rPr = None
    runs = list(p.runs)
    if runs and runs[0]._element.rPr is not None:
        rPr = copy.deepcopy(runs[0]._element.rPr)
    for r in runs:
        r._element.getparent().remove(r._element)
    r = OxmlElement('w:r')
    if rPr is not None:
        r.append(rPr)
    t = OxmlElement('w:t')
    t.text = text
    t.set(qn('xml:space'), 'preserve')
    r.append(t)
    p._p.append(r)

def set_cell_text(cell, text):
    ps = list(cell.paragraphs)
    p0 = ps[0]
    for p in ps[1:]:
        p._element.getparent().remove(p._element)
    rebuild_para(p0, text)

doc = Document(T1)
t = doc.tables[0]

# ---- 1) R2：学院恢复 + 专业填写（网格索引：cells[1]=学院值, cells[6]=专业值） ----
r2 = t.rows[1]
print("R2 cells:", [para_text(c.paragraphs[0]) if c.paragraphs else "" for c in r2.cells][:14])
set_cell_text(r2.cells[1], "计算与信息科学学院")
set_cell_text(r2.cells[6], "计算机科学与技术")

# ---- 2) R4 教师行：重置（网格：cells[2]=张三位, [3]=职称标签, [5]=副教授位, [7]=学位学历标签, [8]=硕士研究生位, [10]=研究方向标签, [11]=网络架构创新位） ----
r4 = t.rows[3]
print("R4 cells:", [para_text(c.paragraphs[0]) if c.paragraphs else "" for c in r4.cells][:14])
set_cell_text(r4.cells[2], PLACEHOLDER)   # 教师姓名值
set_cell_text(r4.cells[3], "职称")        # 标签恢复
set_cell_text(r4.cells[5], PLACEHOLDER)   # 职称值
set_cell_text(r4.cells[7], "学位学历")    # 标签恢复
set_cell_text(r4.cells[8], PLACEHOLDER)   # 学位学历值
set_cell_text(r4.cells[10], "研究方向")   # 标签恢复
set_cell_text(r4.cells[11], PLACEHOLDER)  # 研究方向值

# ---- 3) R9 课题应具备条件：段落级重建 ----
r9 = t.rows[8]
c9 = r9.cells[1]
print("R9 paragraphs:", [para_text(p) for p in c9.paragraphs])
cnt = {"know": 0}
for p in c9.paragraphs:
    txt = para_text(p).strip()
    if txt in ("（1）", "（2）", "（3）"):
        if txt == "（1）" and cnt["know"] == 0:
            rebuild_para(p, "（1）掌握Java面向对象程序设计、数据结构与算法、数据库原理等专业基础知识；")
            cnt["know"] += 1
        elif txt == "（1）" and cnt["know"] >= 1:
            rebuild_para(p, "（1）个人电脑一台（建议16GB以上内存），安装JDK17、Node.js、MySQL8、Redis及微信开发者工具等开发环境；")
        elif txt == "（2）":
            rebuild_para(p, "（2）掌握Spring Boot、MyBatis-Plus等后端框架，以及MySQL、Redis数据库应用技术；")
            cnt["know"] += 1
        elif txt == "（3）":
            rebuild_para(p, "（3）掌握Vue.js等前端开发技术，了解大语言模型基本原理与本地化部署调用方法。")
            cnt["know"] += 1
    elif "工作难度" in txt:
        rebuild_para(p, "（2）工作难度：适中")
    elif "工作任务量" in txt:
        rebuild_para(p, "（3）工作任务量：饱满")
    elif "（2）" in txt and "实验设备" not in txt and "软硬件" not in txt:
        # 第二次出现的（2）段落可能是软硬件第二项（文本为“（2）”+说明残余）
        if "Ollama" not in txt and txt.startswith("（2）"):
            rebuild_para(p, "（2）Ollama本地大模型运行环境（llama3.2等开源模型），具备可联网下载模型与依赖库的网络条件。")

# ---- 4) R10 课题意义：重建 3 个 XXXX 段 ----
r10 = t.rows[9]
c10 = r10.cells[1]
texts = [
    "1.课题意义：本课题针对通用聊天机器人缺乏人格化设定、记忆连续性不足、依赖云端服务存在隐私与成本问题等现状，设计并实现一个基于本地大模型的智能AI陪伴聊天系统。系统通过AI角色与AI精灵双体系提供个性化陪伴，基于“记忆—经历—感受—自我”的灵魂算法构建长期记忆机制，实现“越聊越懂你”的持续陪伴体验；采用Ollama本地化部署保障用户数据隐私、降低使用成本；支持Web端、桌面端与移动端多端使用，并融合好友、群聊、兴趣社群等社交功能。创新点包括：基于Ollama的开源大模型私有化部署方案、融合向量检索与时间衰减的AI记忆机制、AI桌面宠物互动以及三端统一架构。",
    "2.主要研究内容：（1）系统总体架构设计：前后端分离架构，采用Spring Boot 3.2.5、Vue 3、MySQL 8与Redis，支持Web端、Electron桌面端与移动端三端复用；（2）AI角色引擎：角色创建、人设提示词构建、上下文管理、角色一致性对话与流式输出；（3）AI记忆系统：基于向量检索、重要度评分与时间衰减的长期记忆机制；（4）AI精灵与桌面宠物：形象更换、互动反馈与语音对话；（5）社交模块：好友、私聊、群聊、兴趣社群与动态发布；（6）语音交互：语音识别与语音合成；（7）管理员后台：用户管理、数据池管理、公告管理与AI系统管理；（8）安全合规：JWT认证、RBAC权限控制、内容审核与未成年保护；（9）测试体系：单元测试、集成测试与系统测试。",
    "3.主要设计目标：实现功能完整、安全健壮、流畅运行的智能AI陪伴聊天系统，覆盖Web端、桌面端与移动端三端；完成自动化测试用例1000个以上并全部通过，系统全链路稳定运行；交付符合福州理工学院毕业设计规范的设计文档、源代码与毕业论文。",
]
n = 0
for p in c10.paragraphs:
    if para_text(p).strip() == "XXXX" and n < 3:
        rebuild_para(p, texts[n])
        n += 1
print("XXXX replaced:", n)

doc.save(T1)
print("T1 FIXED")
