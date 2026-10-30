# -*- coding: utf-8 -*-
"""开题报告：扩充正文内容 + 插入三张插图（保持批注锚点与正文格式）"""
import copy
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

TARGET = r'C:\Users\郭耀彬\Desktop\gybct\ai-chat-graduation\4.2301132013-XXX-开题报告（填写版）.docx'
FIGS = r'C:\Users\郭耀彬\Desktop\gybct\ai-chat-graduation\_word_tmp\figs'
Q = lambda t: qn('w:' + t)

doc = Document(TARGET)

# ---- 定位正文大格 cell（含"大语言模型自2017年"的表格单元格）----
bigcell = None
for tbl in doc.tables:
    for row in tbl.rows:
        for cell in row.cells:
            txt = cell.text
            if '大语言模型自2017年' in txt or '本课题采用“需求分析' in txt:
                bigcell = cell
                break
        if bigcell is not None:
            break
    if bigcell is not None:
        break
assert bigcell is not None, '未找到正文大格'

def ptext(p):
    return ''.join(r.text or '' for r in p.runs)

def find_para(prefix):
    for p in bigcell.paragraphs:
        if ptext(p).strip().startswith(prefix):
            return p
    raise AssertionError('未找到锚点段: ' + prefix)

# ---- 段落构造工具 ----
def _anchor_el(anchor):
    """兼容 Paragraph 对象与 CT_P 元素"""
    return anchor._p if hasattr(anchor, '_p') else anchor

def make_para_after(anchor, ppr_src, run_rpr_src, text, jc=None, indent=False):
    """在 anchor 段后插入新正文段，继承 ppr_src 段落格式与 run_rpr_src 字符格式"""
    a = _anchor_el(anchor)
    new_p = OxmlElement('w:p')
    # pPr
    if ppr_src is not None:
        pPr = copy.deepcopy(ppr_src)
        # 处理缩进
        ind = pPr.find(Q('ind'))
        if not indent and ind is not None:
            for attr in ('firstLine', 'firstLineChars', 'hanging', 'hangingChars', 'left', 'leftChars'):
                if ind.get(Q(attr)) is not None:
                    del ind.attrib[Q(attr)]
        # 对齐
        if jc is not None:
            j = pPr.find(Q('jc'))
            if j is None:
                j = OxmlElement('w:jc'); pPr.append(j)
            j.set(Q('val'), jc)
        new_p.append(pPr)
    # run
    r = OxmlElement('w:r')
    if run_rpr_src is not None:
        r.append(copy.deepcopy(run_rpr_src))
    t = OxmlElement('w:t')
    t.set(Q('space'), 'preserve')
    t.text = text
    r.append(t)
    new_p.append(r)
    a.addnext(new_p)
    return new_p

def make_pic_para_after(anchor, ppr_src, fig_path, width_cm):
    """在 anchor 元素后插入图片段（居中）"""
    a = _anchor_el(anchor)
    new_p = OxmlElement('w:p')
    if ppr_src is not None:
        pPr = copy.deepcopy(ppr_src)
        ind = pPr.find(Q('ind'))
        if ind is not None:
            for attr in ('firstLine', 'firstLineChars', 'hanging', 'hangingChars', 'left', 'leftChars'):
                if ind.get(Q(attr)) is not None:
                    del ind.attrib[Q(attr)]
        j = pPr.find(Q('jc'))
        if j is None:
            j = OxmlElement('w:jc'); pPr.append(j)
        j.set(Q('val'), 'center')
        new_p.append(pPr)
    a.addnext(new_p)
    para = Paragraph(new_p, bigcell)
    run = para.add_run()
    run.add_picture(fig_path, width=Cm(width_cm))
    return new_p

def make_caption_after(anchor, ppr_src, text):
    """图题段（居中、五号黑体）"""
    a = _anchor_el(anchor)
    new_p = OxmlElement('w:p')
    if ppr_src is not None:
        pPr = copy.deepcopy(ppr_src)
        ind = pPr.find(Q('ind'))
        if ind is not None:
            for attr in ('firstLine', 'firstLineChars', 'hanging', 'hangingChars', 'left', 'leftChars'):
                if ind.get(Q(attr)) is not None:
                    del ind.attrib[Q(attr)]
        j = pPr.find(Q('jc'))
        if j is None:
            j = OxmlElement('w:jc'); pPr.append(j)
        j.set(Q('val'), 'center')
        new_p.append(pPr)
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(Q('eastAsia'), '黑体'); rFonts.set(Q('ascii'), 'Times New Roman'); rFonts.set(Q('hAnsi'), 'Times New Roman')
    rPr.append(rFonts)
    sz = OxmlElement('w:sz'); sz.set(Q('val'), '21'); rPr.append(sz)
    szCs = OxmlElement('w:szCs'); szCs.set(Q('val'), '21'); rPr.append(szCs)
    r.append(rPr)
    t = OxmlElement('w:t'); t.set(Q('space'), 'preserve'); t.text = text
    r.append(t)
    new_p.append(r)
    a.addnext(new_p)
    return new_p

# 获取锚点段及其格式基准
p_bg = find_para('大语言模型自2017年')          # 1.1 背景段
p_r3 = find_para('（3）社交与管理后台')           # 3. 研究内容第(3)段
p_41 = find_para('本课题采用“需求分析')           # 4.1 大段

ppr = p_bg._p.find(Q('pPr'))
rpr = None
for r in p_bg._p.findall(Q('r')):
    rPr = r.find(Q('rPr'))
    if rPr is not None:
        rpr = rPr; break
if rpr is None:
    rpr = OxmlElement('w:rPr')

# ---- 1.1 后新增背景段 ----
add_bg = ('近年来，随着大语言模型技术的快速演进，智能对话应用已由工具型问答向情感陪伴、角色扮演等多元化场景拓展，'
          '用户对可长期陪伴、人格稳定且尊重隐私的AI伙伴需求日益旺盛；与此同时，云端AI服务在数据主权、网络依赖与长期使用成本等方面存在天然短板，'
          '本地化部署的大模型凭借数据不出本机、离线可用、成本可控等优势，成为构建隐私友好型智能陪伴系统的可行技术路线。'
          '这进一步凸显了本课题的研究价值与实践空间。')
make_para_after(p_bg, ppr, rpr, add_bg)

# ---- 3. 研究内容(3)段后：三端协同段 + 图1 ----
add_sd = ('在功能组织上，系统遵循“一端接入、三端协同”的设计思路：Web端面向完整功能管理与深度对话场景，'
          'Electron桌面端提供沉浸式对话与桌面精灵陪伴体验，微信小程序端满足移动场景下的轻量访问需求；'
          '三端共享统一后端服务与数据体系，保证会话、记忆与社交数据的实时一致，其功能分布如图1所示。')
p_sd = make_para_after(p_r3, ppr, rpr, add_sd)
pic1 = make_pic_para_after(p_sd, ppr, FIGS + r'\fig1_sanduan.png', 12.6)
make_caption_after(pic1, ppr, '图1  系统三端功能架构图')

# ---- 4.1 大段后：总体架构段 + 图2 ----
add_arch = ('具体而言，系统总体架构自上而下分为表现层、应用层、服务层与数据层四层：表现层由Web端、Electron桌面端与微信小程序端构成，'
            '通过RESTful API与WebSocket/SSE与后端通信；应用层承载聊天对话、AI角色与AI精灵、记忆引擎、社交互动与系统管理等核心业务模块；'
            '服务层提供大模型推理、语音识别与合成、消息推送等能力，其中大模型统一适配层支持Ollama本地模型的多模型路由与降级切换；'
            '数据层以MySQL 8存储业务数据、Redis缓存热点数据，并由本地模型库承载模型文件。系统总体架构如图2所示。')
p_arch = make_para_after(p_41, ppr, rpr, add_arch)
pic2 = make_pic_para_after(p_arch, ppr, FIGS + r'\fig2_zongti.png', 12.6)
make_caption_after(pic2, ppr, '图2  系统总体架构图')

# ---- 4.1 后：灵魂算法与记忆段 + 图3 ----
add_soul = ('为使AI角色具备持续的人格一致性与成长性，系统将“灵魂算法”思想落地为可执行的记忆与人格演化机制：'
            '以角色当前状态为基座，叠加每一次对话经历、对经历的感受与理解，并结合长期记忆沉淀，共同演化出下一时刻的自我状态；'
            '记忆体系分为短期会话记忆与长期向量记忆两层，通过重要度评分、时间衰减与定期一致性维护，完成记忆的写入、检索与遗忘闭环，其流程如图3所示。')
p_soul = make_para_after(p_arch, ppr, rpr, add_soul)
pic3 = make_pic_para_after(p_soul, ppr, FIGS + r'\fig3_linghun.png', 12.6)
make_caption_after(pic3, ppr, '图3  AI灵魂算法与分层记忆流程图')

doc.save(TARGET)
print('T2 content expanded + 3 figures inserted')
